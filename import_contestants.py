#!/usr/bin/env python3
"""
从Excel文件导入选手名单到contestants.json
使用方法: python3 import_contestants.py <excel文件路径>
"""

import sys
import json
import openpyxl

def import_contestants(excel_file):
    """从Excel文件导入选手数据"""
    try:
        # 加载Excel文件
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        
        print(f"📋 工作表名称: {ws.title}")
        print(f"📊 总行数: {ws.max_row}, 总列数: {ws.max_column}")
        
        # 读取所有行
        contestants = []
        
        # 假设第一行是表头，从第二行开始读取
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):  # 跳过空行
                continue
            
            # 根据实际Excel结构调整索引
            # 常见结构: [组别, 学校, 姓名, 手机号, ...]
            group = row[0] if len(row) > 0 else None
            school = row[1] if len(row) > 1 else None
            name = row[2] if len(row) > 2 else None
            phone = row[3] if len(row) > 3 else None
            
            # 如果姓名为空，尝试其他列
            if not name and len(row) > 0:
                # 尝试找到包含中文的列作为姓名
                for cell in row:
                    if cell and isinstance(cell, str) and any('\u4e00' <= c <= '\u9fff' for c in cell):
                        if len(cell) <= 10:  # 姓名通常较短
                            name = cell
                            break
            
            if name:
                contestant = {
                    "name": str(name).strip(),
                    "phone": str(phone).strip() if phone else "",
                    "school": str(school).strip() if school else "",
                    "group": str(group).strip() if group else ""
                }
                contestants.append(contestant)
        
        print(f"✅ 成功解析 {len(contestants)} 名选手")
        
        # 显示前5名选手作为示例
        print("\n前5名选手示例:")
        for i, c in enumerate(contestants[:5], 1):
            print(f"{i}. {c['name']} | {c['school']} | {c['group']}")
        
        # 保存到JSON文件
        output_file = 'contestants.json'
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump({"contestants": contestants}, f, ensure_ascii=False, indent=4)
        
        print(f"\n✅ 已保存到 {output_file}")
        print(f"📝 总共 {len(contestants)} 名选手")
        
        # 统计组别分布
        groups = {}
        for c in contestants:
            group = c['group']
            groups[group] = groups.get(group, 0) + 1
        
        print("\n📊 组别分布:")
        for group, count in sorted(groups.items()):
            print(f"  - {group}: {count}人")
        
        return contestants
        
    except FileNotFoundError:
        print(f"❌ 错误: 找不到文件 {excel_file}")
        return None
    except Exception as e:
        print(f"❌ 错误: {e}")
        import traceback
        traceback.print_exc()
        return None

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("使用方法: python3 import_contestants.py <excel文件路径>")
        print("\n示例: python3 import_contestants.py contestants.xlsx")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    import_contestants(excel_file)
